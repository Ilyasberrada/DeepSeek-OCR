import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

def add_title_slide(prs):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(9, 25, 46)
    title_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.6), Inches(11), Inches(3.5))
    title_box.fill.solid()
    title_box.fill.fore_color.rgb = RGBColor(31, 81, 112)
    title_box.line.color.rgb = RGBColor(255, 255, 255)
    title_box.shadow.inherit = False
    text_frame = title_box.text_frame
    text_frame.text = "Le travail et la sagesse"
    text_frame.word_wrap = True
    text_frame.paragraphs[0].font.size = Pt(58)
    text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    text_frame.paragraphs[0].font.name = "Montserrat"
    for run in text_frame.paragraphs[0].runs:
        run.font.bold = True
    subtitle_paragraph = text_frame.add_paragraph()
    subtitle_paragraph.text = "Il faut cultiver notre jardin"
    subtitle_paragraph.font.size = Pt(28)
    subtitle_paragraph.font.color.rgb = RGBColor(186, 225, 107)
    subtitle_paragraph.font.name = "Montserrat"
    subtitle_paragraph.level = 1
    subtitle_paragraph.font.italic = True
    circle = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(5.4), Inches(1.35), Inches(2.4), Inches(2.4))
    circle.fill.solid()
    circle.fill.fore_color.rgb = RGBColor(255, 255, 255)
    circle.line.color.rgb = RGBColor(31, 81, 112)
    circle.line.width = Pt(2)
    text_frame = circle.text_frame
    text_frame.text = "3D"
    text_frame.paragraphs[0].font.size = Pt(46)
    text_frame.paragraphs[0].font.bold = True
    text_frame.paragraphs[0].font.color.rgb = RGBColor(31, 81, 112)
    text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

def create_3d_cube(slide, left, top, size, colors):
    front = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, size, size)
    front.fill.solid()
    front.fill.fore_color.rgb = colors['front']
    front.line.color.rgb = RGBColor(255, 255, 255)
    front.line.width = Pt(1.5)
    depth_adj = Inches(0.4)
    offset_adj = Inches(0.28)
    top_shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left + offset_adj, top - depth_adj, size, depth_adj)
    top_shape.fill.solid()
    top_shape.fill.fore_color.rgb = colors['top']
    top_shape.line.color.rgb = RGBColor(255, 255, 255)
    top_shape.line.width = Pt(1.5)
    side = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left + size, top + offset_adj, offset_adj, size)
    side.fill.solid()
    side.fill.fore_color.rgb = colors['side']
    side.line.color.rgb = RGBColor(255, 255, 255)
    side.line.width = Pt(1.5)
    return front

def add_content_slide(prs, title, bullets, cube_colors):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(12, 28, 54)
    cube = create_3d_cube(slide, Inches(0.8), Inches(2.1), Inches(3.0), cube_colors)
    text_frame = cube.text_frame
    text_frame.text = title
    text_frame.paragraphs[0].font.size = Pt(26)
    text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    text_frame.paragraphs[0].font.bold = True
    text_frame.paragraphs[0].font.name = "Raleway"
    text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    body_box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(4.4), Inches(1.3), Inches(8.6), Inches(4.8))
    body_box.fill.solid()
    body_box.fill.fore_color.rgb = RGBColor(20, 45, 79)
    body_box.line.color.rgb = RGBColor(72, 164, 222)
    body_box.line.width = Pt(1.8)
    shadow = body_box.shadow
    shadow.inherit = False
    shadow.color.rgb = RGBColor(0, 0, 0)
    shadow.distance = Pt(15)
    shadow.blur_radius = Pt(20)
    text_frame = body_box.text_frame
    text_frame.text = bullets[0]
    text_frame.word_wrap = True
    first_paragraph = text_frame.paragraphs[0]
    first_paragraph.font.size = Pt(22)
    first_paragraph.font.color.rgb = RGBColor(186, 225, 107)
    first_paragraph.font.bold = True
    first_paragraph.font.name = "Raleway"
    for bullet in bullets[1:]:
        p = text_frame.add_paragraph()
        p.text = bullet
        p.font.size = Pt(18)
        p.font.color.rgb = RGBColor(236, 245, 255)
        p.level = 1
        p.font.name = "Open Sans"
    highlight = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE, Inches(10.6), Inches(5.6), Inches(2.2), Inches(1.6))
    highlight.fill.solid()
    highlight.fill.fore_color.rgb = RGBColor(245, 169, 64)
    highlight.line.color.rgb = RGBColor(255, 255, 255)
    highlight.rotation = -20
    highlight.text_frame.text = "3D"
    highlight.text_frame.paragraphs[0].font.size = Pt(20)
    highlight.text_frame.paragraphs[0].font.bold = True
    highlight.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

def add_quote_slide(prs):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(9, 23, 43)
    ring = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.DONUT, Inches(4.7), Inches(1.3), Inches(4.8), Inches(4.8))
    ring.fill.solid()
    ring.fill.fore_color.rgb = RGBColor(20, 45, 79)
    ring.line.color.rgb = RGBColor(186, 225, 107)
    text_frame = ring.text_frame
    text_frame.text = "Il faut cultiver notre jardin"
    text_frame.word_wrap = True
    text_frame.paragraphs[0].font.size = Pt(26)
    text_frame.paragraphs[0].font.bold = True
    text_frame.paragraphs[0].font.color.rgb = RGBColor(186, 225, 107)
    text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    subtitle = slide.shapes.add_textbox(Inches(1.6), Inches(5.5), Inches(10), Inches(1.3))
    tf = subtitle.text_frame
    tf.text = "— Voltaire, Candide"
    tf.paragraphs[0].font.size = Pt(24)
    tf.paragraphs[0].font.color.rgb = RGBColor(240, 242, 255)
    tf.paragraphs[0].font.italic = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER

def add_conclusion_slide(prs):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(8, 20, 38)
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.PARALLELOGRAM, Inches(0.7), Inches(1.1), Inches(12), Inches(2.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(72, 164, 222)
    bar.line.color.rgb = RGBColor(255, 255, 255)
    bar.text_frame.text = "Sagesse active : Cultiver notre jardin"
    bar.text_frame.paragraphs[0].font.size = Pt(30)
    bar.text_frame.paragraphs[0].font.bold = True
    bar.text_frame.paragraphs[0].font.color.rgb = RGBColor(12, 28, 54)
    bar.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    steps = [
        "Observer la réalité du travail",
        "Identifier nos zones d'influence",
        "Agir avec persévérance et humilité",
        "Récolter les fruits du sens"
    ]
    for idx, step in enumerate(steps):
        textbox = slide.shapes.add_textbox(Inches(1.2 + idx * 3.0), Inches(3.8), Inches(2.4), Inches(1.6))
        tf = textbox.text_frame
        tf.text = step
        tf.paragraphs[0].font.size = Pt(20)
        tf.paragraphs[0].font.color.rgb = RGBColor(240, 242, 255)
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        textbox.line.color.rgb = RGBColor(72, 164, 222)
    underline = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(1), Inches(5.8), Inches(11.6), Inches(0.25))
    underline.fill.solid()
    underline.fill.fore_color.rgb = RGBColor(186, 225, 107)
    underline.line.fill.background()
    footer = slide.shapes.add_textbox(Inches(1.4), Inches(6.4), Inches(10.7), Inches(0.8))
    footer.text_frame.text = "Le travail devient sagesse lorsque nous sommes jardiniers du réel"
    footer.text_frame.paragraphs[0].font.size = Pt(18)
    footer.text_frame.paragraphs[0].font.color.rgb = RGBColor(186, 225, 107)
    footer.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

def main():
    add_title_slide(prs)
    add_content_slide(
        prs,
        "Le travail comme effort créateur",
        [
            "Labeur et créativité construisent nos mondes",
            "Transformer la matière, se transformer soi-même",
            "Apprendre des saisons du métier"
        ],
        {
            'front': RGBColor(31, 81, 112),
            'top': RGBColor(72, 164, 222),
            'side': RGBColor(15, 36, 67)
        }
    )
    add_content_slide(
        prs,
        "Sagesse : le regard intérieur",
        [
            "Réfléchir au sens de nos actions",
            "Mettre en dialogue savoir-faire et savoir-être",
            "Discerner ce qui dépend de nous"
        ],
        {
            'front': RGBColor(45, 96, 128),
            'top': RGBColor(186, 225, 107),
            'side': RGBColor(20, 45, 79)
        }
    )
    add_content_slide(
        prs,
        "Cultiver notre jardin collectif",
        [
            "Créer un environnement fécond pour tous",
            "Collaborer et prendre soin du commun",
            "Récolter ensemble les fruits de l'engagement"
        ],
        {
            'front': RGBColor(66, 130, 158),
            'top': RGBColor(245, 169, 64),
            'side': RGBColor(25, 50, 82)
        }
    )
    add_quote_slide(prs)
    add_conclusion_slide(prs)
    prs.save("le_travail_sagesse_3d.pptx")

if __name__ == "__main__":
    main()
